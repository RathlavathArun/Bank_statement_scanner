"""
Generate sample Excel fixture files for bank statement parser tests.

Each fixture has:
  - 2 title/info rows (bank name, account info)
  - 1 column header row
  - 5-8 sample transactions with realistic 2026-01 dates
  - 1 totals/summary row at the bottom

Run from repo root:
    python apps/api/tests/fixtures/generate_fixtures.py
"""
from __future__ import annotations

import sys
from pathlib import Path

# Ensure openpyxl is importable even when running standalone
try:
    import openpyxl
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

from openpyxl import Workbook

FIXTURE_DIR = Path(__file__).resolve().parent


def _save(wb: Workbook, name: str) -> Path:
    path = FIXTURE_DIR / name
    wb.save(path)
    print(f"  ✓ {path.name}")
    return path


# ─── HDFC Bank ──────────────────────────────────────────────
def generate_hdfc() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Account Statement"

    # 2 title rows
    ws.append(["HDFC BANK Ltd.", "", "", "", "", "", ""])
    ws.append(["Account No: XXXX1234", "Branch: Mumbai Main", "", "", "", "", ""])

    # Header row
    ws.append([
        "Date", "Narration", "Chq./Ref.No.", "Value Dt",
        "Withdrawal Amt.", "Deposit Amt.", "Closing Balance",
    ])

    # Transactions (dd/mm/yy format as per HDFC template)
    txns = [
        ("02/01/26", "UPI/SWIGGY/PAYMENT/123456789", "UPI123456", "02/01/26", 450.00, None, 49550.00),
        ("05/01/26", "NEFT-SALARY-EMPLOYER INC", "NEFT789012", "05/01/26", None, 85000.00, 134550.00),
        ("08/01/26", "UPI/AMAZON/SHOPPING/987654321", "UPI987654", "08/01/26", 2199.00, None, 132351.00),
        ("12/01/26", "ATM-CASH WITHDRAWAL-MUMBAI", "ATM000123", "12/01/26", 10000.00, None, 122351.00),
        ("15/01/26", "IMPS/JOHN DOE/TRANSFER", "IMPS345678", "15/01/26", 5000.00, None, 117351.00),
        ("18/01/26", "NEFT-RENT REFUND-LANDLORD", "NEFT456789", "18/01/26", None, 3000.00, 120351.00),
        ("22/01/26", "POS/BIGBASKET/GROCERIES", "POS567890", "22/01/26", 1875.50, None, 118475.50),
        ("28/01/26", "UPI/PHONEPE/ELECTRICITY BILL", "UPI678901", "28/01/26", 1200.00, None, 117275.50),
    ]
    for txn in txns:
        ws.append(list(txn))

    # Totals row
    ws.append(["", "*** Total ***", "", "", 20724.50, 88000.00, ""])

    return _save(wb, "hdfc_sample.xlsx")


# ─── ICICI Bank ─────────────────────────────────────────────
def generate_icici() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"

    # 2 title rows
    ws.append(["ICICI Bank Limited", "", "", "", "", "", "", ""])
    ws.append(["Account No: XXXX5678", "Statement of Transactions", "", "", "", "", "", ""])

    # Header row
    ws.append([
        "S No.", "Value Date", "Transaction Date", "Cheque Number",
        "Transaction Remarks", "Withdrawal Amount (INR )",
        "Deposit Amount (INR )", "Balance (INR )",
    ])

    txns = [
        (1, "03/01/2026", "03/01/2026", "", "UPI/PAY/ZOMATO/FOOD ORDER", 320.00, None, 74680.00),
        (2, "06/01/2026", "06/01/2026", "CHQ123456", "CHEQUE DEPOSIT FROM CLIENT", None, 25000.00, 99680.00),
        (3, "10/01/2026", "10/01/2026", "", "RTGS-VENDOR PAYMENT-ABC CORP", 15000.00, None, 84680.00),
        (4, "14/01/2026", "14/01/2026", "", "NEFT-BONUS-EMPLOYER PVT LTD", None, 12000.00, 96680.00),
        (5, "20/01/2026", "20/01/2026", "", "UPI/AMAZON/PRIME SUBSCRIPTION", 1499.00, None, 95181.00),
        (6, "25/01/2026", "25/01/2026", "", "IMPS/FRIEND/BIRTHDAY GIFT", 2000.00, None, 93181.00),
    ]
    for txn in txns:
        ws.append(list(txn))

    # Totals row
    ws.append(["", "", "", "", "*** TOTAL ***", 18819.00, 37000.00, ""])

    return _save(wb, "icici_sample.xlsx")


# ─── SBI ────────────────────────────────────────────────────
def generate_sbi() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Account Statement"

    # 2 title rows
    ws.append(["State Bank of India", "", "", "", "", "", ""])
    ws.append(["A/c No: XXXX9012 | Branch: Delhi Connaught Place", "", "", "", "", "", ""])

    # Header row
    ws.append([
        "Txn Date", "Value Date", "Description",
        "Ref No./Cheque No.", "Debit", "Credit", "Balance",
    ])

    txns = [
        ("01/01/2026", "01/01/2026", "UPI/FLIPKART/PURCHASE", "REF001122", 3499.00, None, 46501.00),
        ("04/01/2026", "04/01/2026", "SALARY CREDIT-TCS LTD", "REF334455", None, 65000.00, 111501.00),
        ("09/01/2026", "09/01/2026", "NEFT-RENT PAYMENT-OWNER", "REF556677", 18000.00, None, 93501.00),
        ("13/01/2026", "13/01/2026", "ATM WDL-SBI ATM DELHI", "REF778899", 5000.00, None, 88501.00),
        ("19/01/2026", "19/01/2026", "UPI/PAYTM/MOBILE RECHARGE", "REF990011", 599.00, None, 87902.00),
        ("24/01/2026", "24/01/2026", "FD INTEREST CREDIT", "REF112233", None, 4500.00, 92402.00),
        ("30/01/2026", "30/01/2026", "POS/DMART/HOUSEHOLD", "REF334400", 2150.00, None, 90252.00),
    ]
    for txn in txns:
        ws.append(list(txn))

    # Totals row
    ws.append(["", "", "=== TOTAL ===", "", 29248.00, 69500.00, ""])

    return _save(wb, "sbi_sample.xlsx")


# ─── Axis Bank ──────────────────────────────────────────────
def generate_axis() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"

    # 2 title rows
    ws.append(["Axis Bank Limited", "", "", "", "", ""])
    ws.append(["Account: XXXX3456 | IFSC: UTIB0001234", "", "", "", "", ""])

    # Header row
    ws.append(["Date", "Particulars", "Chq No", "Debit", "Credit", "Balance"])

    txns = [
        ("02-01-2026", "UPI/AMAZON PAY/PAYMENT", "123456", 850.00, None, 59150.00),
        ("07-01-2026", "NEFT FROM EMPLOYER INC", "789012", None, 72000.00, 131150.00),
        ("11-01-2026", "UPI/SWIGGY/FOOD ORDER", "345678", 380.00, None, 130770.00),
        ("16-01-2026", "ATM WITHDRAWAL AXIS ATM", "000001", 5000.00, None, 125770.00),
        ("21-01-2026", "IMPS TO FRIEND ACCOUNT", "567890", 3000.00, None, 122770.00),
        ("26-01-2026", "DIVIDEND CREDIT-MUTUAL FUND", "DIV001", None, 1250.00, 124020.00),
    ]
    for txn in txns:
        ws.append(list(txn))

    # Totals row
    ws.append(["", "** Grand Total **", "", 9230.00, 73250.00, ""])

    return _save(wb, "axis_sample.xlsx")


# ─── Kotak Mahindra Bank ────────────────────────────────────
def generate_kotak() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"

    # 2 title rows
    ws.append(["Kotak Mahindra Bank", "", "", "", "", ""])
    ws.append(["Account No: XXXX7890 | Branch: Bangalore HSR", "", "", "", "", ""])

    # Header row
    ws.append(["Date", "Particulars", "Chq No", "Debit", "Credit", "Balance"])

    txns = [
        ("03-01-2026", "UPI/PHONEPE/RECHARGE", "REF001", 199.00, None, 34801.00),
        ("06-01-2026", "SALARY CREDIT-ABC CORP", "REF002", None, 80000.00, 114801.00),
        ("10-01-2026", "POS/BIGBASKET/GROCERIES", "REF003", 2345.00, None, 112456.00),
        ("14-01-2026", "NEFT OUT/RENT PAYMENT", "REF004", 15000.00, None, 97456.00),
        ("19-01-2026", "UPI/UBER/CAB FARE", "REF005", 450.00, None, 97006.00),
    ]
    for txn in txns:
        ws.append(list(txn))

    # Totals row
    ws.append(["", "--- Total ---", "", 17994.00, 80000.00, ""])

    return _save(wb, "kotak_sample.xlsx")


# ─── Empty workbook (for error-path tests) ──────────────────
def generate_empty() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Empty"
    # Write nothing — just a blank sheet
    return _save(wb, "empty_sample.xlsx")


def main() -> None:
    print("Generating Excel fixtures …")
    generate_hdfc()
    generate_icici()
    generate_sbi()
    generate_axis()
    generate_kotak()
    generate_empty()
    print("Done.")


if __name__ == "__main__":
    main()
