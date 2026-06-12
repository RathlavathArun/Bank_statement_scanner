import io
from collections import defaultdict
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill("solid", fgColor="F5F5F5")
LOW_CONF_FILL = PatternFill("solid", fgColor="FFE0E0")
MID_CONF_FILL = PatternFill("solid", fgColor="FFF3CD")
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def generate_excel(
    transactions: list[dict[str, Any]],
    statement_meta: dict[str, Any],
) -> bytes:
    """Return formatted .xlsx bytes for exported transactions."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"

    headers = [
        "Date",
        "Narration",
        "Debit",
        "Credit",
        "Balance",
        "Payment Mode",
        "Counterparty",
        "Ledger",
        "OCR Confidence",
        "Reviewed",
    ]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for index, tx in enumerate(transactions, start=2):
        amount = Decimal(str(tx.get("amount") or 0))
        debit = float(amount) if tx.get("tx_type") == "DEBIT" else None
        credit = float(amount) if tx.get("tx_type") == "CREDIT" else None
        row = [
            tx.get("date"),
            tx.get("narration"),
            debit,
            credit,
            float(Decimal(str(tx.get("balance")))) if tx.get("balance") not in (None, "") else None,
            tx.get("payment_mode"),
            tx.get("counterparty"),
            tx.get("ledger_name"),
            float(tx["ocr_confidence"]) if tx.get("ocr_confidence") is not None else None,
            "Yes" if tx.get("is_reviewed", True) else "No",
        ]
        sheet.append(row)

        fill = ALT_FILL if index % 2 == 1 else None
        ocr_conf = tx.get("ocr_confidence")
        if ocr_conf is not None:
            ocr_value = float(ocr_conf)
            if ocr_value < 0.70:
                fill = LOW_CONF_FILL
            elif ocr_value < 0.85:
                fill = MID_CONF_FILL

        for cell in sheet[index]:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top")
            if fill is not None:
                cell.fill = fill

    totals_row = len(transactions) + 2
    sheet.cell(row=totals_row, column=1, value="TOTAL")
    sheet.cell(row=totals_row, column=3, value=f"=SUM(C2:C{totals_row - 1})")
    sheet.cell(row=totals_row, column=4, value=f"=SUM(D2:D{totals_row - 1})")
    sheet.cell(row=totals_row, column=5, value=f'=IF(COUNTA(E2:E{totals_row - 1})=0,"",INDEX(E2:E{totals_row - 1},COUNTA(E2:E{totals_row - 1})))')
    for cell in sheet[totals_row]:
        cell.font = Font(bold=True)
        cell.border = THIN_BORDER

    for column in ("C", "D", "E"):
        for row in range(2, totals_row + 1):
            sheet[f"{column}{row}"].number_format = "#,##0.00"

    for column_cells in sheet.columns:
        max_length = 10
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = min(40, max(max_length, len(value) + 2))
        sheet.column_dimensions[column_letter].width = max_length

    sheet.freeze_panes = "A2"

    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Statement"
    summary["B1"] = statement_meta.get("filename") or "statement"
    summary["A2"] = "Bank"
    summary["B2"] = statement_meta.get("bank_id") or "Unknown"
    summary["A3"] = "Date Range"
    dates = [tx.get("date") for tx in transactions if tx.get("date")]
    summary["B3"] = f"{min(dates)} to {max(dates)}" if dates else "N/A"
    summary["A5"] = "Metric"
    summary["B5"] = "Value"

    total_debit = sum(Decimal(str(tx.get("amount") or 0)) for tx in transactions if tx.get("tx_type") == "DEBIT")
    total_credit = sum(Decimal(str(tx.get("amount") or 0)) for tx in transactions if tx.get("tx_type") == "CREDIT")
    summary_rows = [
        ("Total Transactions", len(transactions)),
        ("Total Debit", float(total_debit)),
        ("Total Credit", float(total_credit)),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=6):
        summary.cell(row=row_index, column=1, value=label)
        summary.cell(row=row_index, column=2, value=value)

    payment_mode_totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for tx in transactions:
        payment_mode_totals[tx.get("payment_mode") or "Unknown"] += Decimal(str(tx.get("amount") or 0))

    summary["A10"] = "Payment Mode"
    summary["B10"] = "Total"
    row_index = 11
    for mode, total in sorted(payment_mode_totals.items()):
        summary.cell(row=row_index, column=1, value=mode)
        summary.cell(row=row_index, column=2, value=float(total))
        row_index += 1

    for row in summary.iter_rows():
        for cell in row:
            if cell.row in {1, 5, 10}:
                cell.font = Font(bold=True)
            cell.border = THIN_BORDER

    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 18

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
